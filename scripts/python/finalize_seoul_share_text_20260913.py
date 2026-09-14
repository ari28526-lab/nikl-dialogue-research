"""Apply the user's expanded frequency delivery scope to the recipient documents."""
import json,shutil
from build_seoul_share_package_20260913 import PKG,ROOT,INVENTORY,csv_rows,write_json,make_doc

for r in INVENTORY:
    if r['id']=='seoul_v2':continue
    r['전달']='06_additional_frequency_files의 출처별 ZIP에 확보 빈도 파일 포함'
csv_rows(PKG/'03_frequency_inventory/확보_빈도자료_목록.csv',INVENTORY)
write_json(PKG/'03_frequency_inventory/확보_빈도자료_목록.json',INVENTORY)
readme=PKG/'README.md';s=readme.read_text(encoding='utf-8')
s=s.replace('보유 목록입니다. 목록에 있다는 것이 이번 폴더에 모든 실자료가 있다는 뜻은 아닙니다.','보유 목록입니다. 출처별 실제 빈도 파일은 06_additional_frequency_files의 ZIP에 함께 전달합니다.')
s=s.replace('- 05_provenance:', '- 06_additional_frequency_files: 모두의 말뭉치 빈도 CSV와 JSONL 49쌍, 다층위 LS MP 및 통합판 86쌍, 강범모 김흥규 확보 원표 4개, KoFREN 확보 CSV 4개의 출처별 ZIP입니다. 서울 수정 전 의미표는 포함하지 않습니다.\n- 05_provenance:')
s=s.replace('연구사전 원 DB, 외부 말뭉치 원문 용례 묶음, 강범모 김흥규 원표, KoFREN 원 CSV도 포함하지 않았습니다.','연구사전 원 DB와 외부 말뭉치의 발화 원문 용례 묶음은 포함하지 않았습니다. 강범모 김흥규 확보 원표와 KoFREN 원 CSV는 사용자의 추가 요청에 따라 출처별 ZIP으로 포함했습니다.')
readme.write_text(s,encoding='utf-8')
from docx import Document
p=PKG/'서울코퍼스_구축과_활용_안내.docx';d=Document(p)
replacements={
 '이번에는 보유 목록만 전달.':'빈도 파일과 보유 목록을 함께 전달.',
 '서지와 목록만 전달.':'확보 빈도 원표와 서지 및 목록을 함께 전달.',
 '강범모 김흥규 원표와 KoFREN 원 CSV 및 외부 말뭉치 원문을 함께 재배포하는 묶음은 아니다.':'추가 요청에 따라 강범모 김흥규 확보 원표 4개와 KoFREN 원 CSV 4개 및 다른 말뭉치의 파생 빈도 CSV와 JSONL도 출처별 ZIP에 넣었다. 다른 말뭉치의 발화 원문 용례 묶음은 포함하지 않았다.',
}
paras=list(d.paragraphs)+[p for t in d.tables for row in t.rows for c in row.cells for p in c.paragraphs]
for para in paras:
    s=para.text
    for a,b in replacements.items():s=s.replace(a,b)
    if s!=para.text:
        if para.runs:
            para.runs[0].text=s
            for r in para.runs[1:]:r.text=''
        else:para.add_run(s)
d.save(p)
shutil.copy2(ROOT/'scripts/python/recount_shared_seoul_analysis.py',PKG/'04_code/recount_saved_analysis.py')
print('Expanded scope applied')
