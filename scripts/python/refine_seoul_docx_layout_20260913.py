from build_seoul_share_package_20260913 import PKG
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml.ns import qn
p=PKG/'서울코퍼스_구축과_활용_안내.docx';d=Document(p)
for element in [d.element,d.styles.element]:
    for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
d.styles['Subtitle'].font.italic=False
for index,t in enumerate(d.tables):
    if len(t.columns)==2 and index>0:
        t.columns[0].width=Inches(1.6);t.columns[1].width=Inches(5.4)
        for row in t.rows:row.cells[0].width=Inches(1.6);row.cells[1].width=Inches(5.4)
for para in d.paragraphs:
    if para.text.startswith('저장된 분석만으로 재집계하는 독립 코드'):
        para.text='서울 분석 묶음과 빈도 자료 묶음은 따로 전달할 수 있다. 서울 분석 묶음의 독립 재집계 코드는 외부 API나 음성 없이 작동하며, 제공 빈도표를 지정하면 키별 계수를 대조한다. 최초 구축 코드도 참고 사본으로 제공하지만 개인 사용자명은 자리표시자로 바꾸었다. 최초 분석 재실행에는 원자료, 연구사전 DB, 외부 라이브러리, 경로 설정과 본인의 바른 인증이 필요하다. evidence_refs에 해당하는 사전 근거 원문은 별도 자료다.'
d.save(p)
print('Layout refined')
